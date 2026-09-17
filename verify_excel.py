import openpyxl

wb = openpyxl.load_workbook("CUENTAS MENSUALES CASA 27.xlsx", data_only=False)

print("Sheet names:", wb.sheetnames)

ws = wb["REGISTRO_CUENTAS"]
print("Max row:", ws.max_row)
print("Max column:", ws.max_column)

print("\n--- Header sample ---")
for col_idx in [1, 2, 3, 27, 28, 52, 53, 93, 94, 95, 96, 97]:
    cat = ws.cell(row=3, column=col_idx).value
    sub = ws.cell(row=4, column=col_idx).value
    print(f"Col {col_idx} ({ws.cell(row=4, column=col_idx).coordinate}): Cat='{cat}' | Sub='{sub}'")

print("\n--- Row 5 Sample Formulas ---")
print("Col A (N° ORDEN):", ws["A5"].value)
print("Col B (FECHA):", ws["B5"].value)
print("Col CR (TOTAL):", ws["CR5"].value)

ws_dash = wb["RESUMEN_EJECUTIVO"]
print("\n--- Dashboard Table ---")
for r in range(3, 8):
    row_vals = [ws_dash.cell(row=r, column=c).value for c in range(1, 5)]
    print(f"Row {r}: {row_vals}")

print("\nValidation complete successfully!")
