import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

# Create Workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "CUENTAS MENSUALES"

# Enable Gridlines explicitly
ws.views.sheetView[0].showGridLines = True

# Colors
HEADER_TITLE_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Dark Blue
HEADER_TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="FFFFFF")

CAT_MAJAS_FILL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid") # Navy Blue
CAT_LUCKY_FILL = PatternFill(start_color="548235", end_color="548235", fill_type="solid") # Dark Green
CAT_MERCADO_FILL = PatternFill(start_color="C65911", end_color="C65911", fill_type="solid") # Dark Orange/Rust
CAT_OTROS_FILL = PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid") # Purple
CAT_SUMMARY_FILL = PatternFill(start_color="333333", end_color="333333", fill_type="solid") # Dark Gray

SUBHEADER_MAJAS_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Soft Blue
SUBHEADER_LUCKY_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Soft Green
SUBHEADER_MERCADO_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Soft Orange
SUBHEADER_OTROS_FILL = PatternFill(start_color="E1D5E7", end_color="E1D5E7", fill_type="solid") # Soft Purple
SUBHEADER_SUMMARY_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Soft Gray

WHITE_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
DARK_BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="000000")
REGULAR_FONT = Font(name="Calibri", size=10, bold=False, color="000000")

# Borders
THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
HEADER_BORDER = Border(
    left=Side(style='thin', color='FFFFFF'),
    right=Side(style='thin', color='FFFFFF'),
    top=Side(style='thin', color='FFFFFF'),
    bottom=Side(style='thin', color='FFFFFF')
)

print("Script template ready")
